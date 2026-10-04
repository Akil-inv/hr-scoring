"""
Builds the interview setup workbook: the blank template the app offers for
download, and a realistic sample for testing the upload.

    python3 tools/setup-workbook/build.py

Writes:
    apps/web/public/templates/event-setup-template.xlsx   (served by the app)
    tools/setup-workbook/sample-event.xlsx                (10 weekdays, 9 judges)
    apps/web/public/templates/candidates-template.xlsx    (served by the app)
    tools/setup-workbook/sample-candidates-batch{1,2}.xlsx

The sheet names, column headers and cell formats here are the contract with
the API's workbook reader (apps/api/src/setup-upload/workbook.ts). Change one,
change both.
"""
from datetime import date, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
FONT = "Arial"
REQUIRED = PatternFill("solid", fgColor="1F3864")
OPTIONAL = PatternFill("solid", fgColor="8EA9DB")
EXAMPLE = PatternFill("solid", fgColor="FFF2CC")
RULE = Border(bottom=Side(style="thin", color="BFBFBF"))

# (header, required, width, note)
FIXED_SHEETS = {
    "Event": [
        ("Event name", True, 34, None),
        ("Location", False, 26, None),
        ("Timezone", False, 20, "Where the interviews happen, e.g. Asia/Singapore. Every time in this workbook is in this timezone. Default Asia/Singapore."),
        ("Minimum panel size", False, 12, "Fewest judges an interview needs. Slots with fewer available judges are left without a panel. Default 2."),
        ("Admin emails", False, 40, "Comma-separated. Each must already have a login (Users & roles); they become admins of this event."),
        ("Support question", False, 22, "Asked of every judge as Yes / No, e.g. Support for LAP. Leave blank for none."),
    ],
    "Day template": [
        ("Block", True, 10, "A name for the half-day block, e.g. AM or PM. The Availability sheet uses the same names."),
        ("Start time", True, 11, "24-hour HH:MM. Only on the first row of each block; the rest follow on one after another."),
        ("Item", True, 14, "Interview, Break or Calibration."),
        ("Duration (minutes)", True, 12, None),
    ],
    "Judges": [
        ("Name", True, 22, None),
        ("Email", True, 30, "Must be unique. The Availability sheet refers to judges by email."),
        ("Phone", False, 16, "With country code, e.g. +6591234567"),
        ("Organisation", False, 20, None),
        ("Designation", False, 22, None),
        ("Tier", False, 8, "Optional label: L1, L2, L3, L4, PS or V."),
    ],
    "Rubric": [
        ("Dimension", True, 24, "One row per dimension. Judges rate each 1 to 5 and must comment on each."),
        ("Descriptor", False, 34, "What the dimension is about. Shown to judges."),
        ("Score: 1 (Low)", True, 40, "What a 1 looks like. 2 sits between 1 and 3."),
        ("Score: 3 (Moderate)", True, 40, "What a 3 looks like. 4 sits between 3 and 5."),
        ("Score: 5 (High)", True, 40, "What a 5 looks like."),
        ("Comment required", False, 12, "Y: judges must comment on this dimension. N: comment optional. Default Y."),
        ("Score step", False, 10, "How finely judges can rate this dimension: 1 (whole numbers), 0.5, 0.25 (e.g. 3.75) or 0.1. Default 0.25."),
    ],
    "Criteria": [
        ("Criterion", True, 44, "A category (leave Parent blank) or a row inside one."),
        ("Parent criterion", False, 34, "For a row: the exact name of its category. Blank for a category."),
        ("Max score", True, 10, "Categories must add up to 100. Each category's rows must add up to the category's max."),
        ("Guidance", False, 50, "Shown to judges under the row."),
        ("Comment required", False, 12, "Y if the judge must comment on this row."),
    ],
}

README = [
    ("Interview setup workbook", "title"),
    ("Fill in the sheets and upload this file under Upload setup. The upload checks everything first and shows the schedule it will build. Nothing is saved until you confirm.", None),
    ("", None),
    ("How it works", "head"),
    ("Day template describes one day: each block (AM, PM) is a run of interviews, breaks and calibration with their lengths. Every judging day uses it.", None),
    ("Availability says which judges can sit which block on which date. The system builds every day's slots from the template and seats each judge in the interviews they are available for.", None),
    ("One panel interviews at a time. Interviews with fewer judges than the minimum panel size are left without a panel and shown greyed out.", None),
    ("Candidates are not in this workbook. Once the schedule exists, upload them in batches with the candidates file (Name, Date, Time), as they confirm.", None),
    ("", None),
    ("Availability format (the only accepted entries)", "head"),
    ("Yes = available for the whole block.", None),
    ("No, or blank = not available.", None),
    ("HH:MM-HH:MM in 24-hour time, e.g. 13:00-16:00 = available only within that window. The judge sits only the interviews that fit entirely inside it.", None),
    ("Anything else stops the upload and is listed with the judge and column so it can be corrected.", None),
    ("Column headers on the Availability sheet are the date and block: YYYY-MM-DD AM, YYYY-MM-DD PM. Add one column per date and block; dates with no judges simply stay empty.", None),
    ("", None),
    ("How to fill it in", "head"),
    ("Dark blue header = required column. Light blue header = optional. Hover a header for notes.", None),
    ("Yellow cells are examples. Overwrite or delete them before uploading.", None),
    ("Keep the sheet names and column headers as they are. Times are 24-hour HH:MM.", None),
    ("", None),
    ("Rubric", "head"),
    ("The Rubric sheet holds the interview rubric: each dimension is rated 1 to 5 with a comment, and the candidate's score is the average rating out of 5. It comes filled in with the LAP rubric; edit it for another programme.", None),
    ("Support question on the Event sheet (e.g. Support for LAP) is asked of every judge as Yes / No.", None),
    ("Each Rubric row also sets that dimension's scoring rules. Comment required: Y makes the comment mandatory, N optional (blank = Y). Score step: how finely judges can rate, 1, 0.5, 0.25 or 0.1 (blank = 0.25); with 0.25 a judge can give 3.75.", None),
    ("Criteria is an alternative points rubric (categories adding up to 100). Use either Rubric or Criteria, not both. With neither, the LAP rubric is used.", None),
]

TEMPLATE_DAY = (
    [("AM", time(9, 0), "Interview", 20)]
    + [("AM", None, "Interview", 20)] * 4
    + [("AM", None, "Break", 10)]
    + [("AM", None, "Interview", 20)] * 4
    + [("AM", None, "Calibration", 10)]
    + [("PM", time(14, 0), "Interview", 20)]
    + [("PM", None, "Interview", 20)] * 4
    + [("PM", None, "Break", 10)]
    + [("PM", None, "Interview", 20)] * 4
    + [("PM", None, "Calibration", 10)]
)


LAP_RUBRIC = [
    ("Career Aspirations", "Clarity and ambition regarding future roles and career trajectory",
     "No clear career goals; lacks interest in leadership or generalist roles.",
     "Expresses some interest in leadership but lacks clarity or commitment to generalist path.",
     "Strong aspiration for senior leadership; clearly articulates interest in generalist roles and long-term growth.", "Y", 0.25),
    ("Drive and Motivation", "Energy, initiative, and commitment to personal and organizational goals",
     "Passive attitude; limited examples of initiative or ownership.",
     "Shows moderate drive; some examples of taking initiative or leading efforts.",
     "Highly driven; consistently demonstrates ownership, resilience, and proactive leadership.", "Y", 0.25),
    ("Mobility & Rotation Readiness", "Willingness and preparedness for new roles or rotations",
     "Unwilling to relocate or rotate; prefers stability.",
     "Open to some mobility; hesitant about full rotation model.",
     "Fully open to geographic and functional rotations; embraces diverse experiences.", "Y", 0.25),
    ("Learning Agility & Adaptability", "Ability to learn quickly and adapt to new situations",
     "Resistant to change; struggles with unfamiliar situations.",
     "Some adaptability; has handled change with mixed success.",
     "Highly agile; thrives in new environments and learns quickly from feedback.", "Y", 0.25),
    ("Enterprise Perspective", "Understanding and acting for the broader organization",
     "Narrow focus on own function; lacks cross-functional awareness.",
     "Some awareness of broader business; limited cross-functional experience.",
     "Strong enterprise mindset; demonstrates strategic thinking and cross-functional collaboration.", "Y", 0.25),
]


def style_header(ws, cols):
    for j, (header, required, width, note) in enumerate(cols, 1):
        c = ws.cell(1, j, header + (" *" if required else ""))
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = REQUIRED if required else OPTIONAL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = RULE
        ws.column_dimensions[c.column_letter].width = width
        if note:
            c.comment = Comment(note, "Template")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32


def write_rows(ws, rows, example):
    for r, values in enumerate(rows, 2):
        for j, v in enumerate(values, 1):
            c = ws.cell(r, j, v)
            c.font = Font(name=FONT)
            if example:
                c.fill = EXAMPLE
            if isinstance(v, time):
                c.number_format = "hh:mm"


def choices(ws, rng, options, strict=True):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    if not strict:
        # A dropdown for the common answers that still accepts a typed time window.
        dv.showErrorMessage = False
    dv.add(rng)
    ws.add_data_validation(dv)


def build(data: dict, example: bool) -> Workbook:
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    for i, (line, kind) in enumerate(README, 1):
        c = readme.cell(i, 1, line)
        c.font = Font(name=FONT, bold=kind in ("title", "head"), size=14 if kind == "title" else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 120

    order = ["Event", "Day template", "Judges", "Availability", "Rubric", "Criteria"]
    for name in order:
        ws = wb.create_sheet(name)
        if name == "Availability":
            cols = [("Judge email", True, 30, "Exactly as on the Judges sheet.")]
            cols += [(h, False, 15, "Yes, No, or a window like 13:00-16:00") for h in data["availability_columns"]]
            style_header(ws, cols)
            # Date columns are part of the data, not fixed: style them like required headers.
            for j in range(2, len(cols) + 1):
                ws.cell(1, j).fill = REQUIRED
            write_rows(ws, data["availability"], example)
            last = ws.cell(1, len(cols)).column_letter
            choices(ws, f"B2:{last}500", ["Yes", "No"], strict=False)
            continue
        style_header(ws, FIXED_SHEETS[name])
        write_rows(ws, data.get(name, []), example)
        if name == "Day template":
            choices(ws, "C2:C200", ["Interview", "Break", "Calibration"])
        if name == "Judges":
            choices(ws, "F2:F500", ["L1", "L2", "L3", "L4", "PS", "V"])
        if name == "Criteria":
            choices(ws, "E2:E200", ["Y", "N"])
        if name == "Rubric":
            choices(ws, "F2:F200", ["Y", "N"])
            choices(ws, "G2:G200", ["1", "0.5", "0.25", "0.1"], strict=False)
            # The real rubric, not an example: no yellow, and wrapped to read.
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for c in row:
                    c.fill = PatternFill(fill_type=None)
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    return wb


def template_data() -> dict:
    cols = ["2026-10-19 AM", "2026-10-19 PM", "2026-10-20 AM", "2026-10-20 PM"]
    return {
        "Event": [("UOB Interviews October 2026", "UOB Plaza 1, Singapore", "Asia/Singapore", 2, "admin@example.com", "Support for LAP")],
        "Rubric": LAP_RUBRIC,
        "Day template": TEMPLATE_DAY,
        "Judges": [
            ("Dean Tan", "dean@example.com", "+6591234567", "UOB", "Managing Director", "L2"),
            ("Lawrance Lim", "lawrance@example.com", "+6598765432", "UOB", "Executive Director", "L3"),
            ("Choon Hin Ong", "choonhin@example.com", "+6590001111", "UOB", "Director", "L3"),
        ],
        "availability_columns": cols,
        "availability": [
            ("dean@example.com", "Yes", "Yes", "No", "No"),
            ("lawrance@example.com", "Yes", "13:00-16:00", "No", "Yes"),
            ("choonhin@example.com", "Yes", "Yes", "Yes", "Yes"),
        ],
    }


JUDGES = [
    ("Jack", "jack@example.com"), ("Eric", "eric@example.com"), ("Lay Wah", "laywah@example.com"),
    ("Yung Chee", "yungchee@example.com"), ("Lawrance", "lawrance@example.com"),
    ("Choon Hin", "choonhin@example.com"), ("Hendra", "hendra@example.com"),
    ("Wei Wei", "weiwei@example.com"), ("Dean", "dean@example.com"),
]


def sample_days():
    days, d = [], date(2026, 10, 19)
    while len(days) < 10:
        if d.weekday() < 5:
            days.append(d)
        d += timedelta(days=1)
    return days


# Who can sit which block across the fortnight. Keys are judge email prefixes.
# Thu 22 Oct has nobody (a gap in the schedule); Mon 26 Oct afternoon too.
# Partial windows show judges who can only stay part of a block, and Tue 27
# Oct afternoon drops below two judges after 15:40, so those interviews have
# no panel.
SAMPLE_PLAN = {
    ("2026-10-19", "AM"): {"jack": "Yes", "hendra": "Yes", "weiwei": "Yes"},
    ("2026-10-19", "PM"): {"jack": "Yes", "dean": "Yes", "lawrance": "13:00-16:00"},
    ("2026-10-20", "AM"): {"eric": "Yes", "laywah": "Yes", "yungchee": "Yes"},
    ("2026-10-20", "PM"): {"eric": "Yes", "laywah": "Yes", "choonhin": "14:00-15:00"},
    ("2026-10-21", "AM"): {"dean": "Yes", "lawrance": "Yes", "choonhin": "Yes"},
    ("2026-10-21", "PM"): {"dean": "Yes", "lawrance": "Yes", "choonhin": "Yes"},
    ("2026-10-23", "AM"): {"jack": "Yes", "eric": "Yes"},
    ("2026-10-23", "PM"): {"jack": "Yes", "eric": "Yes", "hendra": "13:00-16:00"},
    ("2026-10-26", "AM"): {"yungchee": "Yes", "weiwei": "Yes", "hendra": "09:00-10:40"},
    ("2026-10-27", "AM"): {"yungchee": "Yes", "weiwei": "Yes", "laywah": "Yes"},
    ("2026-10-27", "PM"): {"dean": "Yes", "hendra": "14:00-15:40"},
    ("2026-10-28", "AM"): {"lawrance": "Yes", "choonhin": "Yes", "dean": "Yes"},
    ("2026-10-28", "PM"): {"lawrance": "Yes", "choonhin": "Yes", "jack": "15:00-17:20"},
    ("2026-10-29", "AM"): {"dean": "Yes", "jack": "Yes", "choonhin": "Yes"},
    ("2026-10-29", "PM"): {"dean": "Yes", "jack": "Yes"},
    ("2026-10-30", "AM"): {"hendra": "Yes", "weiwei": "Yes", "eric": "Yes"},
    ("2026-10-30", "PM"): {"hendra": "Yes", "weiwei": "Yes", "eric": "Yes"},
}


def sample_data() -> dict:
    """Ten weekdays from 19 Oct 2026, nine judges, two gaps, partial windows."""
    cols = [f"{day.isoformat()} {b}" for day in sample_days() for b in ("AM", "PM")]
    availability = []
    for name, email in JUDGES:
        key = email.split("@")[0]
        row = []
        for col in cols:
            d, b = col.split(" ")
            row.append(SAMPLE_PLAN.get((d, b), {}).get(key, "No"))
        availability.append((email, *row))
    return {
        "Event": [("October Graduate Interviews", "UOB Plaza 1, Singapore", "Asia/Singapore", 2, "", "Support for LAP")],
        "Rubric": LAP_RUBRIC,
        "Day template": TEMPLATE_DAY,
        "Judges": [(n, e, f"+659000{i:04d}", "UOB", "Panel member", "L3") for i, (n, e) in enumerate(JUDGES, 1)],
        "availability_columns": cols,
        "availability": availability,
    }


def interview_slots_with_panel(min_panel=2):
    """The interviews the sample schedule gives a panel, as (date, start time), in order."""
    blocks = {}
    for blk, start, kind, minutes in TEMPLATE_DAY:
        items = blocks.setdefault(blk, [])
        at = (start.hour * 60 + start.minute) if start else items[-1][2]
        items.append((kind, at, at + minutes))
    out = []
    for day in sample_days():
        for blk in ("AM", "PM"):
            plan = SAMPLE_PLAN.get((day.isoformat(), blk), {})
            for kind, s, e in blocks[blk]:
                if kind != "Interview":
                    continue
                seated = 0
                for answer in plan.values():
                    if answer == "Yes":
                        seated += 1
                        continue
                    f, t = answer.split("-")
                    fm = int(f[:2]) * 60 + int(f[3:])
                    tm = int(t[:2]) * 60 + int(t[3:])
                    seated += fm <= s and e <= tm
                if seated >= min_panel:
                    out.append((day, time(s // 60, s % 60)))
    return out


FIRST = ["Aisha", "Benjamin", "Chloe", "Daniel", "Elena", "Farid", "Grace", "Hafiz", "Ivy", "Jun Wei",
         "Kavya", "Liam", "Mei Ling", "Nathan", "Olivia", "Priya", "Qistina", "Rahul", "Sarah", "Thanh",
         "Umar", "Vanessa", "Wen Hui", "Xavier", "Yasmin", "Zhi Hao"]
LAST = ["Tan", "Lim", "Rahman", "Ng", "Kumar", "Wong", "Nguyen", "Lee", "Santoso", "Chen", "Ismail", "Goh",
        "Pillai", "Ong", "Wijaya", "Teo", "Hassan", "Chua", "Reyes", "Koh"]


def candidate_names(n):
    names, i = [], 0
    while len(names) < n:
        name = f"{FIRST[i % len(FIRST)]} {LAST[(i * 7 + i // len(FIRST)) % len(LAST)]}"
        if name not in names:
            names.append(name)
        i += 1
    return names


CANDIDATE_COLS = [
    ("Name", True, 30, "The candidate's name. Must be unique in the event; uploading the same name again moves that candidate."),
    ("Date", True, 13, "YYYY-MM-DD, the day of the interview."),
    ("Time", True, 10, "24-hour HH:MM: the start time of an interview slot on the schedule, e.g. 09:20."),
]

CANDIDATE_README = [
    ("Candidates file", "title"),
    ("Upload this under Upload candidates once the schedule exists. Upload as often as you like: each file adds the candidates in it, and moves any already placed whose slot changed. Candidates not in the file are left where they are.", None),
    ("", None),
    ("Name, Date and Time are all required. Time is the start of an interview slot exactly as it appears on the schedule.", None),
    ("A row is refused if the slot doesn't exist, has no panel, already has another candidate, or if the candidate's interview has already started.", None),
]


def build_candidates(rows, example: bool) -> Workbook:
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    for i, (line, kind) in enumerate(CANDIDATE_README, 1):
        c = readme.cell(i, 1, line)
        c.font = Font(name=FONT, bold=kind == "title", size=14 if kind == "title" else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 110
    ws = wb.create_sheet("Candidates")
    style_header(ws, CANDIDATE_COLS)
    write_rows(ws, rows, example)
    for r in range(2, len(rows) + 2):
        ws.cell(r, 2).number_format = "yyyy-mm-dd"
    return wb


def sample_batches():
    """
    Two batches as candidates confirm. Batch 1 fills the first week, leaving
    a few slots open. Batch 2 fills the second week, repeats two people
    unchanged, and moves three first-week candidates into open second-week
    slots.
    """
    slots = interview_slots_with_panel()
    week1 = [s for s in slots if s[0] < date(2026, 10, 26)]
    week2 = [s for s in slots if s[0] >= date(2026, 10, 26)]
    names = candidate_names(len(slots))
    # Leave every 12th first-week slot open, to show gaps being filled later.
    first = [(names[i], d, t) for i, (d, t) in enumerate(week1) if i % 12 != 11]
    open_w2 = week2[-3:]
    second = [(names[len(week1) + i], d, t) for i, (d, t) in enumerate(week2[:-3])]
    second += [first[0], first[1]]                                     # unchanged
    second += [(first[j][0], d, t) for j, (d, t) in zip((2, 3, 4), open_w2)]  # moved
    return first, second


if __name__ == "__main__":
    out = ROOT / "apps/web/public/templates/event-setup-template.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    build(template_data(), example=True).save(out)
    print("wrote", out.relative_to(ROOT))
    sample = ROOT / "tools/setup-workbook/sample-event.xlsx"
    build(sample_data(), example=False).save(sample)
    print("wrote", sample.relative_to(ROOT))
    ctemplate = ROOT / "apps/web/public/templates/candidates-template.xlsx"
    build_candidates([("Candidate One", date(2026, 10, 19), time(9, 0)), ("Candidate Two", date(2026, 10, 19), time(9, 20))], example=True).save(ctemplate)
    print("wrote", ctemplate.relative_to(ROOT))
    first, second = sample_batches()
    for n, rows in ((1, first), (2, second)):
        out = ROOT / f"tools/setup-workbook/sample-candidates-batch{n}.xlsx"
        build_candidates(rows, example=False).save(out)
        print("wrote", out.relative_to(ROOT))
